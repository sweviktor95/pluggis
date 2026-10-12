# -*- coding: utf-8 -*-
"""Biologiplanering på ca 25 veckor. Bygger på momenten i biologi_plan.py (den utökade planen, 77 veckor):
varje delområde får ett urval KÄRNMOMENT (ryms i ca en vecka); resten hamnar under "Om tid / fördjupning".
Kör: python3 planering/biologi_plan25.py   -> biologi.md, biologi-planering.xlsx, biologi-planering.pdf"""
import sys, re, html, collections
sys.path.insert(0, "planering")
import biologi_plan as B
import biologi_extra as X

# ---- Gamla moment indexerade på id ---------------------------------------
OLD = {}
for a in B.P:
    for d in a[3]:
        for i, m in enumerate(d[5], 1):
            if m[0].startswith("Elevlab"): continue
            OLD[f"{d[0]}.{i}"] = dict(titel=m[0], text=m[1], tags=m[2], begrepp=X.BEGREPP[d[0]][i-1], dnr=d[0], area=a[0])

# Ändrade moment (sammanslagna eller omformulerade för att rymmas)
OVR = {
 "1.2.1": dict(titel="Från fråga till plan", text="Göra frågan undersökningsbar; förväntan/hypotes; oberoende, beroende och kontrollerade variabler; rättvist försök.", tags="SU1 S1 S7 F3", begrepp="undersökningsbar fråga, hypotes, oberoende, beroende och kontrollerad variabel, rättvist försök"),
 "1.4.2": dict(titel="Söka information och bedöma källor", text="Söksätt och olika typer av källor; trovärdighet och relevans: avsändare, syfte, reklam och vinklade rubriker.", tags="SU3 S6 F2", begrepp="källa, sökord, trovärdighet, relevans, avsändare, syfte"),
 "4.2.2": dict(titel="Producenter, konsumenter och näringsvävar", text="Roller i ekosystemet; rita och läsa näringskedjor och näringsvävar.", tags="NM4 NM5 F1", begrepp="producent, konsument, nedbrytare, näringskedja, näringsväv, trofinivå"),
 "6.2.1": dict(titel="Matens väg och spjälkning", text="Matspjälkningens organ från mun till tjocktarm; enzymer och upptag i tunntarmen.", tags="KH1 F1", begrepp="matstrupe, magsäck, tunntarm, tjocktarm, enzym, spjälkning, upptag"),
 "8.1.1": dict(titel="Bakterier och virus", text="Byggnad och förökning; nyttiga och skadliga bakterier; virus i värdceller; vad som skiljer dem åt.", tags="KH2 F1", begrepp="bakterie, virus, värdcell, delning, förökning"),
 "8.4.3": dict(titel="Antibiotika och resistens", text="Antibiotika verkar på bakterier, inte virus; resistens uppstår genom naturligt urval; varför antibiotika ska användas rätt.", tags="KH2 NM1 S3 F1", begrepp="antibiotika, resistens, naturligt urval, mutation"),
 "9.1.2": dict(titel="DNA, gener och kromosomer", text="DNA som arvsmassa; gen, kromosom, genom; hur gener styr egenskaper.", tags="NM2 F1", begrepp="DNA, gen, kromosom, genom, cellkärna, protein, egenskap"),
 "11.2.2": dict(titel="Variation och naturligt urval", text="Variation och mutationer som grund; överlevnad och fortplantning; anpassning över generationer.", tags="NM1 NM2 F1", begrepp="variation, mutation, naturligt urval, anpassning, fortplantning"),
 "12.2.2": dict(titel="Sömn och motion", text="Vad sömn och fysisk aktivitet gör för kroppen och psyket; råd och effekter.", tags="KH3 F1", begrepp="sömn, dygnsrytm, kondition, fysisk aktivitet"),
 "12.4.1": dict(titel="Alkohol, tobak, nikotin och droger", text="Effekter och risker; narkotika och lagstiftning.", tags="KH3 F1", begrepp="alkohol, tobak, nikotin, cannabis, narkotika, lagstiftning"),
 "13.2.1": dict(titel="Befruktning och graviditet", text="Från könsceller till foster; graviditetens skeden och hälsa; förlossningen i stora drag.", tags="KH4 F1", begrepp="äggcell, spermie, befruktning, embryo, foster, moderkaka, förlossning"),
 "13.3.1": dict(titel="Preventivmedel och STI", text="Preventivmedel och skydd, nödpreventivmedel; vanliga STI, skydd, test och behandling.", tags="KH4 KH2 F1", begrepp="preventivmedel, kondom, nödpreventivmedel, STI, klamydia, HIV, HPV, provtagning"),
 "13.1.1": dict(titel="Pubertet och könsorgan", text="Hormoner och kroppens förändringar; könsorganens byggnad och funktion; menstruation och sädesutlösning.", tags="KH4 KH1 F1", begrepp="pubertet, könshormon, äggstock, testikel, livmoder, menstruation, sädesutlösning"),
}
def mom(i):
    m = dict(OLD[i]); m.update(OVR.get(i, {})); return m

# ---- Ny plan ----------------------------------------------------------------
# Ca 2,5 h lektion (3 moment) + ca 45 min labb per vecka.
# delområde: (nr, titel, veckor, mål, nivå, [kärna: 3 lektionsmoment], [om tid])
NEW = [
(1, "Biologi som vetenskap", "Hur vi vet det vi vet i biologi: undersöka, värdera, granska och argumentera. Läggs först eftersom allt annat bygger på det.", [
 ("1.1", "Undersöka systematiskt", 1,
  "Eleven kan planera, utföra, dokumentera och värdera en undersökning med en variabel i taget, säkert och med lämplig utrustning.",
  "E: genomför enligt plan och drar enkel slutsats. C: planerar rättvist försök och värderar felkällor. A: motiverar upplägget, föreslår förbättringar och diskuterar hur säker slutsatsen är.",
  ["1.2.1","1.2.3","1.2.4"], ["1.1.1","1.1.2","1.2.2"]),
 ("1.2", "Modeller, källor och argument", 1,
  "Eleven kan förklara att biologisk kunskap prövas och förändras, skilja vad vetenskap kan pröva och granska källor.",
  "E: ger exempel på hur kunskap ändrats och använder flera källor. C: motiverar bedömningen av källorna och kopplar förändring till undersökningar. A: väger källor mot varandra och resonerar om modellers begränsningar.",
  ["1.1.3","1.3.2","1.4.2"], ["1.3.1","1.3.3","1.4.1","1.4.3","1.4.4"]),
]),
(2, "Cellen och livets processer", "Cellen som livets enhet, energin som driver den och vägen från cell till organism.", [
 ("2.1", "Cellen", 1,
  "Eleven kan beskriva och jämföra celltyper och förklara organellernas uppgifter.",
  "E: namnger delar och uppgifter. C: jämför celltyper och kopplar byggnad till funktion. A: förklarar hur cellens delar samverkar i ett förlopp.",
  ["2.1.1","2.1.2","2.1.3"], ["2.1.4","2.1.5"]),
 ("2.2", "Fotosyntes och cellandning", 1,
  "Eleven kan förklara fotosyntes och cellandning och jämföra dem.",
  "E: anger ämnen in/ut. C: jämför processerna och förklarar samband med ljus och koldioxid. A: kopplar processerna till kolets kretslopp och energiflöde.",
  ["2.2.1","2.3.1","2.3.3"], ["2.2.2","2.2.3","2.2.4","2.3.2"]),
 ("2.3", "Organ och organsystem", 0.5,
  "Eleven kan beskriva hur celler bygger upp vävnader, organ och organsystem.",
  "E: ger nivåordningen och exempel. C: förklarar samverkan mellan två organsystem. A: resonerar om hur en störning i ett system påverkar andra.",
  ["2.4.1","2.4.2"], ["2.4.3"]),
]),
(3, "Biologisk mångfald", "Mångfaldens nivåer, hur organismer grupperas, vad som lever omkring oss och varför mångfalden är viktig.", [
 ("3.1", "Mångfald, hot och ekosystemtjänster", 1,
  "Eleven kan beskriva mångfald på tre nivåer, förklara ekosystemtjänster och beskriva hoten mot mångfalden.",
  "E: namnger nivåerna och ger exempel på tjänster och hot. C: kopplar hot till konsekvenser. A: väger olika intressen mot varandra och tar ställning med belägg.",
  ["3.1.1","3.4.1","3.4.2"], ["3.1.2","3.1.3"]),
 ("3.2", "Gruppera och organismvärlden", 1,
  "Eleven kan gruppera organismer systematiskt, bestämma en art med nyckel och beskriva anpassningar.",
  "E: använder en färdig nyckel. C: förklarar principen för gruppering och kopplar byggnad till miljö. A: förklarar anpassningar som resultat av naturligt urval.",
  ["3.2.2","3.2.3","3.3.6"], ["3.2.1","3.2.4","3.3.1","3.3.2","3.3.3","3.3.4","3.3.5"]),
]),
(4, "Ekologi", "Samspel i ekosystem: energi, kretslopp, populationer och resurser, lokalt och globalt.", [
 ("4.1", "Ekosystem och energi", 1,
  "Eleven kan beskriva ekosystem och förklara energiflöde genom näringskedjor och näringsvävar.",
  "E: bygger en näringskedja och namnger faktorer. C: förklarar energiförlust mellan nivåer. A: förutsäger följder av en förändring i en väv.",
  ["4.1.1","4.2.2","4.2.3"], ["4.1.2","4.1.3","4.2.1","4.2.4"]),
 ("4.2", "Materiens kretslopp", 1,
  "Eleven kan beskriva hur kol, kväve och vatten cirkulerar och varför materia inte försvinner.",
  "E: beskriver ett kretslopp. C: kopplar kretslopp till processer. A: resonerar om hur mänsklig påverkan stör kretsloppet.",
  ["4.3.1","4.3.2","4.3.4"], ["4.3.3"]),
 ("4.3", "Populationer och samspel", 1,
  "Eleven kan förklara hur populationer förändras och begränsas av resurser och samspel.",
  "E: beskriver faktorer. C: tolkar diagram över populationsförändring. A: motiverar prognoser och diskuterar osäkerhet.",
  ["4.4.1","4.4.2","4.4.3"], ["4.4.4"]),
]),
(5, "Människan och miljön", "Hur människan påverkar naturen och hur individer och samhället kan främja hållbar utveckling.", [
 ("5.1", "Påverkan på naturen och klimat", 1,
  "Eleven kan förklara hur utsläpp och klimatförändringar påverkar ekosystem.",
  "E: beskriver en påverkan. C: förklarar orsakssamband. A: jämför flera påverkansfaktorer och granskar information om dem.",
  ["5.1.3","5.2.1","5.2.2"], ["5.1.1","5.1.2","5.1.4","5.2.3"]),
 ("5.2", "Hållbar utveckling och ställningstagande", 1,
  "Eleven kan förklara hållbar utveckling och argumentera för val på individ- och samhällsnivå.",
  "E: framför enkelt argument med viss underbyggnad. C: bemöter motargument. A: väger etiska och naturvetenskapliga skäl mot varandra.",
  ["5.3.1","5.3.2","5.3.3"], ["5.3.4"]),
]),
(6, "Kroppens system: mat, andning och blod", "Hur kroppens organsystem tar upp och transporterar det cellerna behöver, och hur de samverkar.", [
 ("6.1", "Mat och matspjälkning", 1,
  "Eleven kan beskriva näringsämnenas uppgifter, följa maten genom kroppen och förklara hur kosten påverkar hälsan.",
  "E: ger uppgifter och organ. C: förklarar enzymer, upptag och kostråd. A: kopplar byggnad till funktion och bedömer en måltids näringsinnehåll.",
  ["6.1.1","6.2.1","6.3.1"], ["6.1.2","6.1.3","6.2.2","6.2.3","6.3.2","6.3.3"]),
 ("6.2", "Andning, blod och samverkan", 1,
  "Eleven kan förklara gasutbyte och blodomlopp och hur organsystem samverkar vid ansträngning.",
  "E: beskriver väg och delar. C: förklarar syretransport och samverkan. A: förutsäger och förklarar kroppens reaktion på ansträngning.",
  ["7.1.1","7.2.1","7.2.3"], ["7.1.2","7.2.2","7.3.1","7.3.2"]),
]),
(7, "Infektioner och immunförsvar", "Bakterier och virus, smitta, kroppens försvar, vaccin och antibiotika.", [
 ("7.1", "Bakterier, virus och smitta", 1,
  "Eleven kan jämföra bakterier och virus, förklara smittvägar och hur infektioner förebyggs.",
  "E: ger skillnader och smittvägar. C: förklarar förökning och förebyggande åtgärder. A: analyserar smittspridning och åtgärder på samhällsnivå.",
  ["8.1.1","8.2.1","8.2.2"], ["8.1.2"]),
 ("7.2", "Immunförsvar, vaccin och antibiotika", 1,
  "Eleven kan beskriva immunförsvaret, förklara vaccin och antibiotika och resistens som evolution.",
  "E: beskriver vad vaccin och antibiotika gör. C: förklarar flockimmunitet och resistens. A: granskar påståenden och argumenterar för åtgärder.",
  ["8.3.2","8.4.1","8.4.3"], ["8.3.1","8.3.3","8.4.2","8.4.4"]),
]),
(8, "Arv, genetik och genteknik", "Arvsmassan, hur egenskaper ärvs, samspelet mellan arv och miljö samt genteknikens möjligheter och etik.", [
 ("8.1", "DNA och nedärvning", 1,
  "Eleven kan beskriva DNA, gener och kromosomer och förklara nedärvning med korsningsschema.",
  "E: namnger begreppen och tolkar enkla scheman. C: räknar sannolikheter. A: förklarar sambandet gen–protein–egenskap och tolkar stamträd.",
  ["9.1.2","9.3.1","9.3.2"], ["9.1.1","9.1.3","9.2.1","9.2.2","9.3.3"]),
 ("8.2", "Arv och miljö, mutationer och genteknik", 1,
  "Eleven kan förklara arv och miljö, beskriva gentekniska metoder och argumentera om möjligheter, risker och etik.",
  "E: ger exempel och framför enkelt argument. C: förklarar samverkan och bemöter argument med belägg. A: väger flera perspektiv mot varandra.",
  ["9.4.1","10.1.2","10.3.2"], ["9.4.2","10.1.1","10.2.1","10.2.2","10.2.3","10.3.1"]),
]),
(9, "Evolution och livets historia", "Hur livet uppstod och förändras, mekanismerna och bevisen. Evolutionen bär även övriga områden.", [
 ("9.1", "Livets historia och bevis", 1,
  "Eleven kan beskriva livets historia och naturvetenskapliga förklaringar till livets uppkomst och värdera bevis för evolution.",
  "E: återger en tidslinje och ger bevis. C: skiljer belagt från hypotes. A: väger bevis mot varandra.",
  ["11.1.1","11.1.3","11.3.1"], ["11.1.2","11.3.2"]),
 ("9.2", "Evolutionens mekanismer", 1,
  "Eleven kan förklara naturligt urval, variation och artbildning och skilja naturvetenskapliga från andra sätt att förklara.",
  "E: beskriver urval. C: förklarar med exempel och skiljer ut vad som kan prövas. A: tillämpar mekanismerna på nya fall och resonerar sakligt om olika sätt att förklara.",
  ["11.2.2","11.2.3","11.4.3"], ["11.2.1","11.4.1","11.4.2"]),
]),
(10, "Hälsa och levnadsvanor", "Hur levnadsvanor och levnadsförhållanden påverkar hälsan och vad som kan göras på individ- och samhällsnivå.", [
 ("10.1", "Stress, sömn och psykisk hälsa", 1,
  "Eleven kan förklara stress, sömn och motion och hur levnadsvanor och levnadsförhållanden påverkar hälsan.",
  "E: beskriver stressreaktion och nytta av sömn och motion. C: kopplar levnadsförhållanden till hälsa. A: analyserar åtgärder på individ- och samhällsnivå.",
  ["12.2.1","12.2.2","12.2.3"], ["12.1.1","12.1.2","12.2.4","12.3.1","12.3.2"]),
 ("10.2", "Beroendeframkallande medel", 1,
  "Eleven kan förklara hur beroende uppstår och vad som kan begränsa problemen.",
  "E: beskriver risker. C: förklarar beroende. A: argumenterar om åtgärder på samhällsnivå.",
  ["12.4.1","12.4.3","12.4.4"], ["12.4.2"]),
]),
(11, "Pubertet, sexualitet och relationer", "Kroppens förändringar, reproduktion, skydd, relationer och samtycke.", [
 ("11.1", "Pubertet, reproduktion, preventivmedel och STI", 1,
  "Eleven kan beskriva pubertet, befruktning och graviditet och förklara preventivmedel och STI.",
  "E: ger exempel. C: jämför metoder och förklarar förloppet. A: motiverar val med belägg.",
  ["13.1.1","13.2.1","13.3.1"], ["13.1.2","13.2.2","13.2.3","13.3.2","13.3.3"]),
 ("11.2", "Relationer, identitet och samtycke", 1,
  "Eleven kan resonera om sexualitet, identitet, relationer och samtycke.",
  "E: förklarar samtycke. C: resonerar om ansvar och ömsesidighet. A: analyserar situationer med flera perspektiv.",
  ["13.4.1","13.4.2","13.4.3"], ["13.4.4"]),
]),
(12, "Repetition inför nationella proven", "Binder ihop begrepp, undersökningar och argumentation.", [
 ("12.1", "Repetition inför nationella proven", 1,
  "Eleven kan förklara samband mellan områden, planera och värdera en undersökning och argumentera.",
  "E: återger och planerar enkelt. C: förklarar samband och värderar upplägg. A: kopplar flera områden i ett resonemang och argumenterar med belägg.",
  ["14.1.1","14.2.1","14.2.2"], ["14.1.2","14.3.1"]),
]),
]

# Labbpass (ca 45 min) per delområde: (titel, undersökningsbar fråga, upplägg, säkerhet, koder, begrepp)
LAB25 = {
 "1.1": ("Reaktionstid som rättvist försök", "Påverkar övning reaktionstiden?", "Varierar: antal försök. Mäter: fallsträcka med fallande linjal. Kontrollerar: hand, startläge, mätare. Planeras i lektionerna, utförs och dokumenteras i passet.", "Låg risk. Resultaten behandlas i klassen, inte som personliga värden.", "SU1 S7 S8 F3", "variabel, mätning, medelvärde, felkälla"),
 "1.2": ("Granska ett påstående", "Går ett vanligt hälsopåstående att belägga?", "Välj ett påstående, sök två källor och bedöm trovärdighet och relevans enligt en checklista.", "Välj neutrala påståenden i klassen; inga personuppgifter.", "SU3 S6 F2", "påstående, källa, trovärdighet, relevans"),
 "2.1": ("Mikroskopera celler", "Hur ser lök- och växtceller ut i mikroskop?", "Gör ett preparat, rita och märk cellen, jämför med en bild av djurcell.", "Mikroskop och objektglas hanteras varsamt; skärverktyg bara med lärare.", "KH1 SU1 S8 F3", "mikroskop, preparat, cellvägg, cellkärna"),
 "2.2": ("Jäst och koldioxid", "Hur påverkar temperaturen jästens koldioxidbildning?", "Varierar: temperatur. Mäter: ballongens omfång eller skumhöjd efter fast tid.", "Använd ljummet, inte hett vatten; tvätta händerna.", "NM5 SU1 S8 F3", "cellandning, jäsning, koldioxid, kontrollförsök"),
 "3.1": ("Artinventering i en provruta", "Hur många sorters växter finns i en kvadratmeter?", "Räkna olika sorter i provrutor på två platser, jämför i tabell.", "Smaka aldrig på växter eller svampar; följ allemansrätten.", "NM6 SU1 S8 F3", "inventering, provruta, artantal"),
 "3.2": ("Artbestämning med nyckel", "Vilka arter är det? Bestäm med nyckel.", "Bestäm löv eller barr med bestämningsnyckel och dokumentera med bild.", "Samla bara vanliga arter; handtvätt.", "NM1 SU1 S8 F3", "bestämningsnyckel, kännetecken, art"),
 "4.1": ("Energi i en näringskedja – modellspel", "Hur mycket energi når toppen av en kedja?", "Modell med kort och marker: energi minskar mellan nivåerna; noteras i tabell.", "Låg risk.", "NM5 SU1 F3", "energiförlust, trofinivå, modell"),
 "4.2": ("Nedbrytning i jord", "Hur snabbt bryts olika material ned?", "Starta ett försök med olika material i jord; avläs vid nästa pass.", "Handskar och handtvätt; samla in materialet efteråt.", "NM5 SU1 S8 F3", "nedbrytning, nedbrytare, kontroll"),
 "4.3": ("Rovdjur–bytesmodell", "Hur påverkar tillgången på föda en population?", "Modellspel med kort; ändra födomängd och notera populationen i tabell och diagram.", "Låg risk.", "NM4 SU1 F3", "population, begränsande faktor, bärkraft"),
 "5.1": ("Växthuseffekten – modell", "Hur påverkar en 'atmosfär' temperaturen?", "Två flaskor med termometer under lampa, den ena med lock. Modellen förenklar verkligheten.", "Lampa blir het; vuxen hanterar.", "NM5 NM6 SU1 S8 F3", "växthuseffekt, modell, mätning"),
 "5.2": ("Argumentera i en miljöfråga", "Hur kan vi väga naturvetenskapliga och etiska skäl?", "Strukturerad diskussion med roller; samla belägg och bemöt motargument.", "Fokusera på argument, inte personer.", "SU4 S4 S5 F2", "argument, motargument, ståndpunkt"),
 "6.1": ("Enzymförsök", "Hur påverkar temperatur nedbrytning av stärkelse?", "Varierar: temperatur. Mäter: tid tills jodtestet inte längre ger färg.", "Jodlösning fläckar och kan irritera; skyddsglasögon. Smaka inte på något.", "KH1 SU1 S8 F3", "enzym, stärkelse, jodtest"),
 "6.2": ("Puls och andning vid belastning", "Hur förändras puls och andning efter ansträngning?", "Mät vila och efter belastning; sammanställ i diagram.", "Frivilligt; anpassa belastningen så att alla kan delta.", "KH1 KH3 SU1 S8 F3", "puls, andningsfrekvens, belastning"),
 "7.1": ("Smittspridningsmodell med handtvätt", "Hur snabbt sprids smitta, och hur påverkar handtvätt?", "Modellmaterial (glitter eller fluorescerande pulver) med och utan handtvätt; räkna 'smittade'.", "Inga odlingar från kroppen eller ytor.", "KH2 SU1 F3", "smittväg, hygien, modell"),
 "7.2": ("Granska ett vaccinpåstående", "Går ett påstående om vaccin att belägga?", "Sök två källor, bedöm trovärdighet och relevans, sammanfatta.", "Välj påståenden i klassen; inga personuppgifter.", "KH2 SU3 S6 F2", "vaccin, källa, belägg"),
 "8.1": ("DNA ur frukt", "Går DNA att göra synligt?", "Jordgubbe med diskmedel, salt och kall sprit; jämför med och utan diskmedel.", "Sprit är brandfarlig och ska inte intas; skyddsglasögon.", "NM2 SU1 S8 F3", "DNA, extraktion"),
 "8.2": ("Etisk bedömning av genteknik", "Vilka argument finns för och emot en genteknisk användning?", "Strukturerad diskussion med roller, belägg och motargument.", "Fokusera på argument, inte personer.", "NM3 SU4 S4 S5 F2", "möjlighet, risk, etik, argument"),
 "9.1": ("Jämför skelett och bygg släktträd", "Vad visar likheter i byggnad om släktskap?", "Jämför ben i framben hos olika däggdjur med digitala bilder; rita ett enkelt släktträd.", "Låg risk.", "NM1 SU2 S8 F3", "homologa organ, släktskap, släktträd"),
 "9.2": ("Modellförsök om urval", "Hur påverkar kamouflage chansen att bli upptäckt?", "Färgade bytespunkter på olika bakgrund; mät andel som hittas.", "Låg risk.", "NM1 SU1 SU2 F3", "naturligt urval, kamouflage, modell"),
 "10.1": ("Sömn och vanor: statistik", "Vad säger statistik om sömn bland unga?", "Läs och tolka ett diagram från en myndighet; formulera en slutsats och en fråga.", "Ingen enkät med personlig information i klassen.", "KH3 SU1 SU3 S8 F3", "statistik, diagram, slutsats"),
 "10.2": ("Läs och tolka statistik om tobak och alkohol", "Vad visar statistiken om användning bland unga?", "Tolka ett diagram från en myndighet; ange vad diagrammet visar och inte visar.", "Inga personliga berättelser i klassen.", "KH3 SU1 SU3 S8 F3", "statistik, diagram, urval"),
 "11.1": ("Faktagranska information om skydd", "Hur kan man bedöma information om preventivmedel och STI?", "Granska två kvalitetssäkrade källor och en vilseledande; jämför.", "Känsligt ämne: frågor ställs anonymt; inga personliga frågor.", "KH4 SU3 S6 F2", "källa, relevans, trovärdighet"),
 "11.2": ("Samtycke – fallövning", "Hur bedömer vi samtycke i olika situationer?", "Diskutera färdiga fall i grupp: vad är samtycke, gränser, var får man hjälp.", "Fiktiva fall; alla får avstå från att dela.", "KH4 S5 SU4 F2", "samtycke, gräns, ömsesidighet"),
 "12.1": ("Provlik mini-undersökning", "Kan du planera och värdera en undersökning på egen hand?", "En färdig frågeställning: planera, genomför enkelt, värdera.", "Låg risk.", "SU1 S7 F3", "variabel, felkälla, slutsats"),
}


# ---- Kontroll --------------------------------------------------------------
def kontroll():
    used = collections.Counter(); seen = collections.Counter(); area_f = collections.defaultdict(set); s3 = set()
    def tag(area, tags):
        for t in B.expand(tags):
            seen[t] += 1
            if t.startswith("F"): area_f[area].add(t)
            if t == "S3": s3.add(area)
    for a in NEW:
        for d in a[3]:
            for i in d[5] + d[6]: used[i] += 1
            for i in d[5]: tag(a[0], mom(i)["tags"])
            if d[0] in LAB25: tag(a[0], LAB25[d[0]][4])
    missing_ids = sorted(set(OLD) - set(used)); dup = [i for i, c in used.items() if c > 1]; unk = [i for i in used if i not in OLD]
    ej = sorted(B.CODES - set(seen))
    weeks = sum(d[2] for a in NEW for d in a[3]); nk = sum(len(d[5]) for a in NEW for d in a[3]); ne = sum(len(d[6]) for a in NEW for d in a[3])
    nl = sum(1 for a in NEW for d in a[3] if d[0] in LAB25)
    bad = [d[0] for a in NEW for d in a[3] if d[2] >= 1 and len(d[5]) != 3]
    return seen, area_f, s3, missing_ids, dup, unk, ej, weeks, nk, ne, nl, bad

def kl(ks, pat): return [f"{k} {X.SHORT[k]}" for k in ks if re.match(pat, k)]
PS, PC, PF = r"S\d", r"(NM|KH|SU)", r"F"
def vecka(v): return f"{v:g}".replace(".", ",")

def modeller(d):
    return list(dict.fromkeys(X.MODELLER[dn] for dn in dict.fromkeys(OLD[i]["dnr"] for i in d[5]) if dn in X.MODELLER))

def main():
    seen, area_f, s3, missing_ids, dup, unk, ej, weeks, nk, ne, nl, bad = kontroll()
    print("veckor", weeks, "kärnmoment", nk, "om tid", ne, "labbpass", nl)
    print("saknade id:", missing_ids, "dubbletter:", dup, "okända:", unk, "kodtäckning saknas:", ej, "ej 3 moment:", bad)
    print("S3 i områden:", sorted(s3))
    for a in NEW: print(a[0], sorted(area_f[a[0]]), vecka(sum(d[2] for d in a[3])), "v")
    if missing_ids or dup or unk or ej or bad: sys.exit(1)
    skriv_md(seen, s3, weeks, nk, ne, nl); skriv_xlsx(); skriv_pdf()

# ---- Markdown ---------------------------------------------------------------
def skriv_md(seen, s3, weeks, nk, ne, nl):
    o = []; w = o.append
    w("# Biologi åk 7–9 – planering på ca 25 veckor (Lgr22)\n")
    w("Status: förslag, väntar på granskning. Genereras av `planering/biologi_plan25.py` – ändra där, inte här.\n")
    w(f"Planen är kortad till ca {vecka(weeks)} veckor: {len(NEW)} områden, {sum(len(a[3]) for a in NEW)} delområden (ca en vecka vardera), {nk} lektionsmoment och {nl} labbpass. Det som inte ryms står under *Om tid / fördjupning* ({ne} moment) och kan användas i fördjupningshäftena eller om tiden räcker.\n")
    w("Tidsram: ca 2,5 timmar lektion per vecka (tre moment à ca 50 minuter) plus ca 45 minuter labb per vecka. Ett delområde tar ca en vecka: tre lektionsmoment och ett labbpass.\n")
    w("Koder: **S1–S8** = ämnets syfte, **NM/KH/SU** = centralt innehåll, **F1–F3** = förmågorna. Syftesdelar som följer av förmåga eller innehåll är härledda: F1→S2, F3→S7, SU3→S6, SU4→S4. Lektionsmoment och labbpass tillsammans täcker hela det centrala innehållet och alla delar av syftet (kontrollerat av skriptet).\n")
    w("## Översikt\n"); w("| Område | Delområden | Veckor |"); w("|---|---|---|")
    for a in NEW: w(f"| {a[0]}. {a[1]} | {', '.join(d[1] for d in a[3])} | {vecka(sum(d[2] for d in a[3]))} |")
    w(f"| **Totalt** | | **{vecka(weeks)}** |\n")
    w("## Planen\n")
    for a in NEW:
        w(f"## {a[0]}. {a[1]}\n"); w(f"*{a[2]}*\n")
        for d in a[3]:
            w(f"### {d[0]} {d[1]} ({vecka(d[2])} v)\n"); w(f"**Mål:** {d[3]}\n"); w(f"**Nivåer (E→A):** {d[4]}\n")
            mod = modeller(d)
            if mod: w("**Förklaringsmodeller:** " + "; ".join(mod) + "\n")
            w("| Moment | Innehåll | Begrepp | Syfte | Centralt innehåll | Förmåga |"); w("|---|---|---|---|---|---|")
            for n, i in enumerate(d[5], 1):
                m = mom(i); ks = B.expand(m["tags"])
                w(f"| **{d[0]}.{n} {m['titel']}** | {m['text']} | {m['begrepp']} | {'; '.join(kl(ks,PS))} | {'; '.join(kl(ks,PC))} | {'; '.join(kl(ks,PF))} |")
            if d[0] in LAB25:
                t, fr, var, sak, tg, beg = LAB25[d[0]]; ks = B.expand(tg)
                w(f"| 🧪 **Labb (ca 45 min): {t}** | {fr} {var} Säkerhet: {sak} | {beg} | {'; '.join(kl(ks,PS))} | {'; '.join(kl(ks,PC))} | {'; '.join(kl(ks,PF))} |")
            w("")
            if d[6]: w("**Om tid / fördjupning:** " + "; ".join(mom(i)["titel"] for i in d[6]) + "\n")
    w("## Öppna frågor\n")
    w("1. Stämmer antagandet att de 25 veckorna gäller hela biologin åk 7–9?")
    w("2. Sexualitet och relationer (område 11) har bara 2 veckor. Är det rimligt?")
    w("3. Evolutionen (område 9) ligger sent. Ska den flyttas tidigare?")
    w("4. Delområde 2.3 är en halv vecka och har inget labbpass.")
    w("5. Sajten har två nivåer under ämnet idag; tre nivåer kräver en ändring av ämnesträdet.")
    open("planering/biologi.md", "w", encoding="utf-8").write("\n".join(o) + "\n")

# ---- Excel -------------------------------------------------------------------
def skriv_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    F = Font(name="Arial", size=10); FB = Font(name="Arial", size=10, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="3E7D5A")
    th = Side(style="thin", color="CCCCCC"); BD = Border(top=th, bottom=th, left=th, right=th)
    def sheet(ws, head, rows, widths, freeze="A2"):
        ws.append(head)
        for r in rows: ws.append(r)
        for c in ws[1]: c.font = FB; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="center")
        for row in ws.iter_rows(min_row=2):
            for c in row: c.font = F; c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = BD
        for i, wd in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = wd
        ws.freeze_panes = freeze; ws.auto_filter.ref = ws.dimensions
    wb = Workbook(); mrows = []; drows = []; arows = []; lrows = []
    for a in NEW:
        arows.append([a[0], a[1], a[2], sum(d[2] for d in a[3]), len(a[3])])
        for d in a[3]:
            drows.append([a[0], a[1], d[0], d[1], d[2], d[3], d[4], "; ".join(modeller(d)), "; ".join(mom(i)["titel"] for i in d[6])])
            for n, i in enumerate(d[5], 1):
                m = mom(i); ks = B.expand(m["tags"])
                mrows.append([a[0], a[1], d[0], d[1], d[2], "Lektionsmoment", f"{d[0]}.{n}", m["titel"], m["text"], m["begrepp"], "; ".join(kl(ks, PS)), "; ".join(kl(ks, PC)), "; ".join(kl(ks, PF))])
            if d[0] in LAB25:
                t, fr, var, sak, tg, beg = LAB25[d[0]]; ks = B.expand(tg)
                mrows.append([a[0], a[1], d[0], d[1], d[2], "Labb (45 min)", f"{d[0]}.L", t, f"{fr} {var}", beg, "; ".join(kl(ks, PS)), "; ".join(kl(ks, PC)), "; ".join(kl(ks, PF))])
                lrows.append([d[0], d[1], t, fr, var, sak])
            for i in d[6]:
                m = mom(i); ks = B.expand(m["tags"])
                mrows.append([a[0], a[1], d[0], d[1], d[2], "Om tid / fördjupning", "", m["titel"], m["text"], m["begrepp"], "; ".join(kl(ks, PS)), "; ".join(kl(ks, PC)), "; ".join(kl(ks, PF))])
    ws = wb.active; ws.title = "Moment"
    sheet(ws, ["Område nr", "Område", "Delområde nr", "Delområde", "Veckor", "Typ", "Moment nr", "Moment", "Innehåll i texten", "Begrepp", "Syfte", "Centralt innehåll", "Förmåga"],
          mrows, [8, 24, 9, 28, 8, 16, 9, 32, 55, 36, 36, 40, 22], "I2")
    sheet(wb.create_sheet("Delområden"), ["Område nr", "Område", "Delområde nr", "Delområde", "Veckor", "Mål", "Nivåer (E→A)", "Förklaringsmodeller", "Om tid / fördjupning"], drows, [8, 24, 9, 30, 8, 50, 60, 45, 50])
    sheet(wb.create_sheet("Områden"), ["Nr", "Område", "Idé", "Veckor", "Delområden"], arows, [6, 32, 70, 8, 10])
    sheet(wb.create_sheet("Labbar (45 min)"), ["Delområde nr", "Delområde", "Labb", "Undersökningsbar fråga", "Upplägg", "Säkerhet och anmärkning"], lrows, [10, 30, 32, 40, 60, 50])
    seen = kontroll()[0]
    sheet(wb.create_sheet("Täckning"), ["Kod", "Innehåll", "Antal lektioner/labbar"], [[k, v, seen[k]] for k, v in list(B.SYFTE.items()) + list(B.CI.items()) + list(B.FORM.items())], [8, 110, 14])
    wo = wb.create_sheet("Om planen", 0)
    for r in ["Biologi åk 7–9 – planering på ca 25 veckor (Lgr22). Förslag, väntar på granskning.", "Nivåer: område → delområde (ca 1 vecka) → moment.", "Tidsram: ca 2,5 h lektion (tre moment) + ca 45 min labb per vecka.", "Syftesdelar härleds: F1→S2, F3→S7, SU3→S6, SU4→S4. S1, S3, S5, S8 är satta för hand.", "Begrepp, labbar och veckor är förslag, inte granskade mot läromedel eller skolans säkerhetsregler.", "Källa: planering/biologi_plan25.py i repot (ändra där, inte här)."]:
        wo.append([r])
    for row in wo.iter_rows():
        for c in row: c.font = F
    wo["A1"].font = Font(name="Arial", size=12, bold=True); wo.column_dimensions["A"].width = 120
    wb.save("planering/biologi-planering.xlsx")

# ---- PDF ---------------------------------------------------------------------
def skriv_pdf():
    e = html.escape
    css = """@page{size:A4;margin:16mm 14mm 16mm 14mm;@bottom-center{content:counter(page);font:9pt Arial;color:#666}}
body{font:9.5pt/1.35 Arial,'DejaVu Sans',sans-serif;color:#1c2b22}
h1{font-size:20pt;margin:0 0 4pt;color:#2f6b4c}.sub{color:#555;margin-bottom:10pt}
div.pg{break-after:page;break-inside:avoid}div.pg:last-child{break-after:auto}
h2{font-size:15pt;color:#fff;background:#3e7d5a;padding:5pt 8pt;border-radius:4pt;margin:0 0 3pt}
.idea{font-style:italic;color:#444;margin:0 0 8pt}
h3{font-size:11.5pt;color:#2f6b4c;border-bottom:1.5px solid #3e7d5a;margin:10pt 0 3pt;padding-bottom:1pt}
h3 span{font-weight:normal;color:#555;font-size:9.5pt}
.mal,.niv,.mod{margin:2pt 0}
table{width:100%;border-collapse:collapse;margin:3pt 0 4pt;table-layout:fixed}
th{background:#e6f1ea;text-align:left;font-size:8.5pt;padding:3pt 4pt;border:1px solid #c5d3c9}
td{vertical-align:top;padding:3pt 4pt;border:1px solid #d5ddd7;font-size:8.8pt}
tr{break-inside:avoid}td.m b{display:block}td.k div{margin-bottom:1.5pt}
tr.lab td{background:#f3f9f5}
.tid{margin:5pt 0;padding:4pt 8pt;background:#f4f6f5;border-radius:4pt;font-size:8.8pt}
.leg{font-size:9pt;background:#f4f6f5;border-radius:4pt;padding:6pt 9pt;margin:6pt 0}.leg li{margin:1pt 0}"""
    o = ['<!doctype html><meta charset="utf-8"><style>' + css + "</style>"]
    weeks = kontroll()[7]
    o.append(f'<div class="pg"><h1>Biologi åk 7–9 – planering</h1><div class="sub">Ca {vecka(weeks)} veckor: område → delområde (ca en vecka) → moment, med begrepp, förklaringsmodeller, labbpass och koppling till kursplanen (Lgr22). Förslag, väntar på granskning.</div>')
    o.append('<div class="leg"><b>Så läser du planen.</b><ul style="margin:3pt 0 0 14pt;padding:0"><li><b>Tidsram:</b> ca 2,5 timmar lektion (tre moment) plus ca 45 minuter labb per vecka.</li><li><b>Område</b> = ett kapitel. <b>Delområde</b> = ca en veckas arbete. <b>Moment</b> = en lektion: vad som tas upp i texten.</li><li><b>Koppling</b> visar vilken del av ämnets <b>syfte</b> (S), vilket <b>centralt innehåll</b> (NM, KH, SU) och vilken <b>förmåga</b> (F1 begrepp/modeller, F2 granska/ta ställning, F3 undersöka) momentet bygger på.</li><li>Nivåraden visar vad som skiljer godtagbara (E), goda (C) och mycket goda (A) kunskaper i delområdet.</li><li><b>Om tid / fördjupning</b> = moment som inte ryms i veckan.</li><li>🧪 = labbpass (ca 45 min).</li></ul></div>')
    o.append('<table><tr><th>Område</th><th>Delområden</th><th style="width:12%">Veckor</th></tr>')
    for a in NEW: o.append(f'<tr><td><b>{a[0]}. {e(a[1])}</b></td><td>{e(", ".join(d[1] for d in a[3]))}</td><td>{vecka(sum(d[2] for d in a[3]))}</td></tr>')
    o.append(f'<tr><td><b>Totalt</b></td><td></td><td><b>{vecka(weeks)}</b></td></tr></table></div>')
    def rad(titel_html, text, beg, ks, cls=""):
        sy = kl(ks, PS); ce = kl(ks, PC); fm = kl(ks, PF); k = ""
        if sy: k += "<div><b>Syfte:</b> " + e("; ".join(sy)) + "</div>"
        if ce: k += "<div><b>Innehåll:</b> " + e("; ".join(ce)) + "</div>"
        k += "<div><b>Förmåga:</b> " + e("; ".join(fm)) + "</div>"
        return f'<tr class="{cls}"><td class="m"><b>{titel_html}</b>{e(text)}</td><td>{e(beg)}</td><td class="k">{k}</td></tr>'
    for a in NEW:
        for di, d in enumerate(a[3]):
            o.append('<div class="pg">')
            if di == 0: o.append(f'<h2>{a[0]}. {e(a[1])}</h2><div class="idea">{e(a[2])} ({vecka(sum(x[2] for x in a[3]))} veckor)</div>')
            else: o.append(f'<div class="idea" style="margin:0 0 4pt;color:#3e7d5a"><b>{a[0]}. {e(a[1])}</b></div>')
            o.append(f'<h3>{d[0]} {e(d[1])} <span>({vecka(d[2])} v)</span></h3><div class="mal"><b>Mål:</b> {e(d[3])}</div><div class="niv"><b>Nivåer:</b> {e(d[4])}</div>')
            mod = modeller(d)
            if mod: o.append('<div class="mod"><b>Förklaringsmodeller:</b> ' + e("; ".join(mod)) + "</div>")
            o.append('<table><colgroup><col style="width:43%"><col style="width:22%"><col style="width:35%"></colgroup><tr><th>Moment och innehåll</th><th>Begrepp</th><th>Koppling till kursplanen</th></tr>')
            for n, i in enumerate(d[5], 1):
                m = mom(i); o.append(rad(f"{d[0]}.{n} {e(m['titel'])}", m["text"], m["begrepp"], B.expand(m["tags"])))
            if d[0] in LAB25:
                t, fr, var, sak, tg, beg = LAB25[d[0]]
                o.append(rad(f"🧪 Labb (ca 45 min): {e(t)}", f"{fr} {var} Säkerhet: {sak}", beg, B.expand(tg), "lab"))
            o.append("</table>")
            if d[6]: o.append('<div class="tid"><b>Om tid / fördjupning:</b> ' + e("; ".join(mom(i)["titel"] for i in d[6])) + "</div>")
            o.append("</div>")
    open("/tmp/plan.html", "w", encoding="utf-8").write("\n".join(o))
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        b = pw.chromium.launch(executable_path="/opt/pw-browsers/chromium"); pg = b.new_page()
        pg.goto("file:///tmp/plan.html"); pg.pdf(path="planering/biologi-planering.pdf", prefer_css_page_size=True, print_background=True); b.close()

if __name__ == "__main__":
    main()
